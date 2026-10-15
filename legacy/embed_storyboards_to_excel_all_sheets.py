#!/usr/bin/env python3
from __future__ import annotations

import argparse
import os
import re
import sys
from copy import copy
from pathlib import Path

try:
    from PIL import Image as PILImage
except ImportError:
    print("Missing dependency: Pillow. Install with: pip install pillow", file=sys.stderr)
    raise

try:
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Missing dependency: openpyxl. Install with: pip install openpyxl", file=sys.stderr)
    raise


SHOT_RE = re.compile(r"(?:^|[^A-Za-z0-9])S(\d{1,2})(?:[^0-9]|$)", re.IGNORECASE)
DEFAULT_SHEETS = (
    "01_47镜头总表",
    "02_真实拍摄层",
    "03_AI生成层",
    "04_MotionGraphics层",
)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description=(
            "Compress storyboard PNGs to same-resolution JPEG quality 65 and embed previews "
            "into the master sheet plus Real / AI / Motion Graphics sub-sheets."
        )
    )
    p.add_argument(
        "--root",
        type=Path,
        default=Path(__file__).resolve().parent,
        help="Project directory. Default: directory containing this script",
    )
    p.add_argument(
        "--png-dir",
        default="generated_storyboards",
        help="PNG source directory under root",
    )
    p.add_argument(
        "--jpg-dir",
        default="generated_storyboards_jpeg65",
        help="JPEG output directory under root",
    )
    p.add_argument(
        "--excel",
        type=Path,
        default=None,
        help="Input xlsx. Default: auto-detect the base ZYUbuntu workbook under root",
    )
    p.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output xlsx. Default: input stem + _含AI分镜_四表.xlsx",
    )
    p.add_argument(
        "--quality",
        type=int,
        default=65,
        help="JPEG quality, default 65",
    )
    p.add_argument(
        "--preview-width",
        type=int,
        default=320,
        help="Displayed image width in Excel pixels, default 320",
    )
    p.add_argument(
        "--preview-height",
        type=int,
        default=180,
        help="Displayed image max height in Excel pixels, default 180",
    )
    p.add_argument(
        "--variant",
        default="v01",
        help="Prefer filenames containing this variant, default v01. Use '' to disable",
    )
    p.add_argument(
        "--alpha-bg",
        choices=["black", "white"],
        default="black",
        help="Background used if a PNG has transparency",
    )
    p.add_argument(
        "--overwrite-jpeg",
        action="store_true",
        help="Rebuild JPEGs even if an up-to-date JPEG already exists",
    )
    p.add_argument(
        "--in-place",
        action="store_true",
        help="Overwrite source workbook instead of writing a new workbook",
    )
    p.add_argument(
        "--sheets",
        default=",".join(DEFAULT_SHEETS),
        help="Comma-separated sheet names to receive previews",
    )
    return p.parse_args()


def find_excel(root: Path) -> Path:
    candidates: list[tuple[int, float, Path]] = []
    for p in root.glob("*.xlsx"):
        if p.name.startswith("~$"):
            continue
        # Prefer the pristine source workbook so reruns do not accumulate images.
        if "含AI分镜" in p.stem:
            continue
        score = 0
        if "影视制作脚本" in p.name:
            score += 10
        if "ZYUbuntu" in p.name:
            score += 5
        candidates.append((score, p.stat().st_mtime, p))

    if not candidates:
        raise FileNotFoundError(
            f"No base .xlsx found in {root}. Pass it explicitly with --excel PATH."
        )
    candidates.sort(reverse=True)
    return candidates[0][2]


def shot_no_from_name(name: str) -> int | None:
    m = SHOT_RE.search(name)
    if not m:
        m = re.match(r"S(\d{1,2})", name, re.IGNORECASE)
    if not m:
        return None
    n = int(m.group(1))
    return n if 1 <= n <= 99 else None


def collect_pngs(png_dir: Path, variant: str) -> dict[int, Path]:
    by_shot: dict[int, list[Path]] = {}
    for p in sorted(png_dir.rglob("*.png")):
        n = shot_no_from_name(p.name)
        if n is not None:
            by_shot.setdefault(n, []).append(p)

    chosen: dict[int, Path] = {}
    for n, files in by_shot.items():
        preferred = [p for p in files if variant and variant.lower() in p.name.lower()]
        pool = preferred or files
        # If several files match the same shot/variant, use the newest.
        chosen[n] = max(pool, key=lambda p: p.stat().st_mtime)
    return chosen


def flatten_to_rgb(im: PILImage.Image, bg: str) -> PILImage.Image:
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        rgba = im.convert("RGBA")
        color = (0, 0, 0, 255) if bg == "black" else (255, 255, 255, 255)
        base = PILImage.new("RGBA", rgba.size, color)
        base.alpha_composite(rgba)
        return base.convert("RGB")
    return im.convert("RGB")


def convert_png_to_jpeg(
    src: Path,
    dst: Path,
    quality: int,
    bg: str,
    overwrite: bool,
) -> tuple[int, int]:
    dst.parent.mkdir(parents=True, exist_ok=True)
    if dst.exists() and not overwrite and dst.stat().st_mtime >= src.stat().st_mtime:
        with PILImage.open(dst) as im:
            return im.size

    with PILImage.open(src) as im:
        original_size = im.size
        rgb = flatten_to_rgb(im, bg)
        # Deliberately do NOT resize. The JPEG keeps the PNG's original pixel dimensions.
        rgb.save(
            dst,
            format="JPEG",
            quality=quality,
            optimize=True,
            progressive=True,
            subsampling=2,  # 4:2:0; large file-size reduction, suitable for storyboard previews
            exif=b"",
        )
    return original_size


def find_header_row_and_col(ws, label: str, max_rows: int = 12) -> tuple[int, int]:
    for row in ws.iter_rows(min_row=1, max_row=min(max_rows, ws.max_row)):
        for cell in row:
            if str(cell.value).strip() == label:
                return cell.row, cell.column
    raise ValueError(f"Header '{label}' not found in sheet {ws.title}")


def copy_header_style(src, dst) -> None:
    if src.has_style:
        dst._style = copy(src._style)
    if src.font:
        dst.font = copy(src.font)
    if src.fill:
        dst.fill = copy(src.fill)
    if src.border:
        dst.border = copy(src.border)
    if src.alignment:
        dst.alignment = copy(src.alignment)
    if src.number_format:
        dst.number_format = src.number_format
    if src.protection:
        dst.protection = copy(src.protection)


def remove_existing_images_in_columns(ws, columns: set[int]) -> None:
    """Remove only images anchored in our preview columns, keeping unrelated sheet images."""
    kept = []
    for img in getattr(ws, "_images", []):
        anchor = getattr(img, "anchor", None)
        col = None
        if hasattr(anchor, "_from"):
            col = anchor._from.col + 1
        elif isinstance(anchor, str):
            m = re.match(r"([A-Z]+)", anchor)
            if m:
                val = 0
                for ch in m.group(1):
                    val = val * 26 + ord(ch) - 64
                col = val
        if col not in columns:
            kept.append(img)
    ws._images = kept


def prepare_preview_columns(ws, header_row: int) -> tuple[int, int]:
    existing_headers: dict[str, int] = {}
    for cell in ws[header_row]:
        if cell.value is not None:
            existing_headers[str(cell.value).strip()] = cell.column

    preview_col = existing_headers.get("AI分镜预览")
    path_col = existing_headers.get("JPEG路径")

    next_col = ws.max_column + 1
    if preview_col is None:
        preview_col = next_col
        next_col += 1
    if path_col is None:
        path_col = next_col

    style_source_col = max(1, min(preview_col - 1, ws.max_column))
    style_src = ws.cell(header_row, style_source_col)

    preview_header = ws.cell(header_row, preview_col, "AI分镜预览")
    path_header = ws.cell(header_row, path_col, "JPEG路径")
    copy_header_style(style_src, preview_header)
    copy_header_style(style_src, path_header)

    ws.column_dimensions[get_column_letter(preview_col)].width = 46
    ws.column_dimensions[get_column_letter(path_col)].width = 34
    return preview_col, path_col


def build_shot_row_map(ws, header_row: int, shot_col: int) -> dict[int, int]:
    out: dict[int, int] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        value = ws.cell(r, shot_col).value
        try:
            n = int(value)
        except (TypeError, ValueError):
            continue
        if 1 <= n <= 99:
            out[n] = r
    return out


def embed_sheet(
    ws,
    jpg_by_shot: dict[int, Path],
    output: Path,
    preview_width: int,
    preview_height: int,
) -> tuple[int, list[int]]:
    header_row, shot_col = find_header_row_and_col(ws, "镜号")
    preview_col, path_col = prepare_preview_columns(ws, header_row)

    # Makes rerunning against an already-embedded workbook safer.
    remove_existing_images_in_columns(ws, {preview_col})

    row_for_shot = build_shot_row_map(ws, header_row, shot_col)
    embedded = 0
    missing: list[int] = []

    for shot, row in sorted(row_for_shot.items()):
        jpg = jpg_by_shot.get(shot)
        if jpg is None:
            missing.append(shot)
            continue

        img = XLImage(str(jpg))
        # This changes only the displayed size in Excel. The embedded JPEG retains full resolution.
        scale = min(preview_width / img.width, preview_height / img.height, 1.0)
        img.width = max(1, int(img.width * scale))
        img.height = max(1, int(img.height * scale))
        ws.add_image(img, f"{get_column_letter(preview_col)}{row}")

        desired_pt = img.height * 0.75 + 8
        current = ws.row_dimensions[row].height or 15
        ws.row_dimensions[row].height = max(current, desired_pt)

        rel = os.path.relpath(jpg, output.parent)
        path_cell = ws.cell(row, path_col, rel)
        path_cell.hyperlink = rel
        path_cell.style = "Hyperlink"
        path_cell.alignment = Alignment(vertical="center", wrap_text=True)

        embedded += 1

    # Extend autofilter if present.
    if ws.auto_filter and ws.auto_filter.ref:
        try:
            start = ws.auto_filter.ref.split(":")[0]
            ws.auto_filter.ref = (
                f"{start}:{get_column_letter(max(preview_col, path_col))}{ws.max_row}"
            )
        except Exception:
            pass

    return embedded, missing


def main() -> int:
    args = parse_args()
    if not 1 <= args.quality <= 95:
        raise ValueError("--quality must be between 1 and 95")

    root = args.root.expanduser().resolve()
    png_dir = (root / args.png_dir).resolve()
    jpg_dir = (root / args.jpg_dir).resolve()

    if not png_dir.exists():
        raise FileNotFoundError(f"PNG directory does not exist: {png_dir}")

    excel = args.excel.expanduser().resolve() if args.excel else find_excel(root)
    if not excel.exists():
        raise FileNotFoundError(excel)

    if args.in_place:
        output = excel
    elif args.output:
        output = args.output.expanduser().resolve()
    else:
        output = excel.with_name(excel.stem + "_含AI分镜_四表.xlsx")

    sheets = [x.strip() for x in args.sheets.split(",") if x.strip()]
    if not sheets:
        raise ValueError("No sheets specified")

    selected = collect_pngs(png_dir, args.variant)
    if not selected:
        raise RuntimeError(f"No S01-style PNG files found under {png_dir}")

    print(f"Project      : {root}")
    print(f"PNG source   : {png_dir}")
    print(f"JPEG output  : {jpg_dir}")
    print(f"Excel input  : {excel}")
    print(f"Excel output : {output}")
    print(f"JPEG quality : {args.quality}, same pixel dimensions")
    print(f"Sheets       : {', '.join(sheets)}")
    print(f"Shots found  : {len(selected)}")

    jpg_by_shot: dict[int, Path] = {}
    png_total = 0
    jpg_total = 0

    print("\nCompressing PNG -> JPEG...")
    for shot, src in sorted(selected.items()):
        dst = jpg_dir / f"S{shot:02d}.jpg"
        w, h = convert_png_to_jpeg(
            src, dst, args.quality, args.alpha_bg, args.overwrite_jpeg
        )
        jpg_by_shot[shot] = dst
        png_total += src.stat().st_size
        jpg_total += dst.stat().st_size
        print(
            f"  S{shot:02d}: {src.name} -> {dst.name}  "
            f"{w}x{h}  {dst.stat().st_size / 1024:.0f} KB"
        )

    wb = load_workbook(excel)
    missing_sheets = [name for name in sheets if name not in wb.sheetnames]
    if missing_sheets:
        raise KeyError(
            "Sheet(s) not found: " + ", ".join(missing_sheets) +
            f". Available: {wb.sheetnames}"
        )

    total_embedded = 0
    sheet_stats: list[tuple[str, int, list[int]]] = []
    print("\nEmbedding previews...")
    for name in sheets:
        ws = wb[name]
        count, missing = embed_sheet(
            ws,
            jpg_by_shot,
            output,
            args.preview_width,
            args.preview_height,
        )
        total_embedded += count
        sheet_stats.append((name, count, missing))
        print(f"  {name}: {count} image(s)")

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    ratio = (jpg_total / png_total * 100.0) if png_total else 0.0
    output_size = output.stat().st_size if output.exists() else 0

    print("\nDone.")
    print(f"JPEG total        : {jpg_total / 1024 / 1024:.2f} MB")
    print(f"PNG total         : {png_total / 1024 / 1024:.2f} MB")
    print(f"JPEG/PNG          : {ratio:.1f}%")
    print(f"Total embeddings  : {total_embedded}")
    print(f"Workbook size     : {output_size / 1024 / 1024:.2f} MB")
    print(f"Workbook          : {output}")

    for name, _, missing in sheet_stats:
        if missing:
            print(
                f"WARNING {name}: no JPEG found for "
                + ", ".join(f"S{x:02d}" for x in missing)
            )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())